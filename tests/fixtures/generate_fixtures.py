import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx

fixtures_dir = os.path.dirname(os.path.abspath(__file__))
os.makedirs(fixtures_dir, exist_ok=True)

# 1. Generate synthetic PDF for Spike A
pdf_path = os.path.join(fixtures_dir, "synthetic_spike_a.pdf")
doc = SimpleDocTemplate(pdf_path, pagesize=letter)
styles = getSampleStyleSheet()

story = []
story.append(Paragraph("<b>MASTER SERVICES AGREEMENT</b>", styles["Title"]))
story.append(Spacer(1, 12))
story.append(Paragraph("This Master Services Agreement is entered into between Client and Supplier.", styles["Normal"]))
story.append(Spacer(1, 12))
story.append(Paragraph("<b>Section 1. Scope of Work</b>", styles["Heading2"]))
story.append(Paragraph("1. Supplier shall provide software development and consulting services.", styles["Normal"]))
story.append(Paragraph("2. Deliverables shall be provided in accordance with the project schedule.", styles["Normal"]))
story.append(Spacer(1, 12))
story.append(Paragraph("<b>Section 2. Limitation of Liability</b>", styles["Heading2"]))
story.append(Paragraph("Supplier's aggregate liability under this Agreement shall not exceed AED 100,000.", styles["Normal"]))
story.append(Spacer(1, 12))
story.append(Paragraph("<b>Section 3. Payment Schedule</b>", styles["Heading2"]))

table_data = [
    ["Milestone", "Deliverable", "Amount (AED)"],
    ["1", "Initial Architecture & Spike", "25,000"],
    ["2", "Core Ingestion Engine", "35,000"],
    ["3", "Final Acceptance & Deployment", "40,000"],
]
table = Table(table_data, colWidths=[80, 260, 100])
table.setStyle(TableStyle([
    ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
    ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
    ('ALIGN', (0,0), (-1,-1), 'LEFT'),
    ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
    ('BOTTOMPADDING', (0,0), (-1,0), 6),
    ('GRID', (0,0), (-1,-1), 1, colors.grey),
]))
story.append(table)

doc.build(story)
print(f"Generated PDF: {pdf_path}")

# 2. Generate synthetic DOCX for Spike B
docx_path = os.path.join(fixtures_dir, "synthetic_spike_b.docx")
doc_docx = docx.Document()

doc_docx.add_heading("MASTER SERVICES AGREEMENT", level=0)
p1 = doc_docx.add_paragraph("This Master Services Agreement is entered into between Client and Supplier.")

h1 = doc_docx.add_heading("Section 1. Scope of Work", level=1)
doc_docx.add_paragraph("1. Supplier shall provide software development and consulting services.")
doc_docx.add_paragraph("2. Deliverables shall be provided in accordance with the project schedule.")

h2 = doc_docx.add_heading("Section 2. Limitation of Liability", level=1)
p_liability = doc_docx.add_paragraph("Supplier's aggregate liability under this Agreement shall not exceed AED 100,000.")

h3 = doc_docx.add_heading("Section 3. Payment Schedule", level=1)
t = doc_docx.add_table(rows=1, cols=3)
hdr_cells = t.rows[0].cells
hdr_cells[0].text = "Milestone"
hdr_cells[1].text = "Deliverable"
hdr_cells[2].text = "Amount (AED)"

milestones = [
    ("1", "Initial Architecture & Spike", "25,000"),
    ("2", "Core Ingestion Engine", "35,000"),
    ("3", "Final Acceptance & Deployment", "40,000"),
]
for m, d, a in milestones:
    row_cells = t.add_row().cells
    row_cells[0].text = m
    row_cells[1].text = d
    row_cells[2].text = a

doc_docx.save(docx_path)
print(f"Generated DOCX: {docx_path}")

# 3. Generate 5-page synthetic PDF with page 5 blank
from reportlab.platypus import PageBreak
from reportlab.graphics.shapes import Drawing, Rect

five_page_path = os.path.join(fixtures_dir, "synthetic_5page_with_blank.pdf")
doc_5p = SimpleDocTemplate(five_page_path, pagesize=letter)
story_5p = []

# Page 1: paragraph
story_5p.append(Paragraph("<b>MASTER SERVICES AGREEMENT - PAGE 1</b>", styles["Title"]))
story_5p.append(Paragraph("This Agreement is made on January 15, 2026 by and between Enterprise Client and Prime Supplier.", styles["Normal"]))
story_5p.append(Paragraph("Both parties mutually agree to the terms, conditions, and schedules set forth herein.", styles["Normal"]))
story_5p.append(PageBreak())

# Page 2: clause
story_5p.append(Paragraph("<b>Clause 1: Confidentiality Obligations</b>", styles["Heading2"]))
story_5p.append(Paragraph("Each party agrees to maintain in strict confidence all proprietary technical, financial, and business information received from the disclosing party.", styles["Normal"]))
story_5p.append(Paragraph("The receiving party shall not disclose Confidential Information to any third party without prior written consent.", styles["Normal"]))
story_5p.append(PageBreak())

# Page 3: clause
story_5p.append(Paragraph("<b>Clause 2: Limitation of Liability</b>", styles["Heading2"]))
story_5p.append(Paragraph("In no event shall either party's aggregate liability arising out of or related to this Agreement exceed AED 500,000.", styles["Normal"]))
story_5p.append(Paragraph("Neither party shall be liable for indirect, incidental, special, or consequential damages.", styles["Normal"]))
story_5p.append(PageBreak())

# Page 4: paragraph
story_5p.append(Paragraph("<b>Section 3: Governing Law and Dispute Resolution</b>", styles["Heading2"]))
story_5p.append(Paragraph("This Agreement shall be governed by and construed in accordance with English law and the jurisdiction of the London Commercial Court.", styles["Normal"]))
story_5p.append(PageBreak())

# Page 5: completely blank (just a tiny transparent spacer or empty drawing so page 5 exists with 0 text)
d = Drawing(10, 10)
d.add(Rect(0, 0, 10, 10, fillColor=colors.white, strokeColor=colors.white))
story_5p.append(d)

doc_5p.build(story_5p)
print(f"Generated 5-Page PDF: {five_page_path}")

# 4. Generate Scanned (Image-Only) PDF with no extractable text
from reportlab.pdfgen import canvas
scanned_path = os.path.join(fixtures_dir, "synthetic_scanned_image.pdf")
c = canvas.Canvas(scanned_path, pagesize=letter)
# Draw an image-like rect / drawing with no text operator
c.setFillColor(colors.lightgrey)
c.rect(50, 50, 500, 700, fill=1, stroke=1)
c.showPage()
c.save()
print(f"Generated Scanned PDF: {scanned_path}")

